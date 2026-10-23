#!/usr/bin/env python3
"""
Documentos · Word engine (separate from MotherBrain's generator).

  list <path>             → JSON [{"i": n, "text": "..."}]  every paragraph of the body, in document order
  hastags <path>          → "1" if the document has {{TAGS}}, else "0"
  apply <path> <out>      ← JSON ops on stdin: [{"op": "replace"|"insert_after"|"delete", "i": n, "text": "..."}]
  fill <path> <out>       ← JSON {"TAG": "value" | ["1st", "2nd"…]} on stdin: exact {{TAG}} filling (tags without a value stay)
  spans <path> <out>      ← JSON [{"i": n, "start": a, "end": b, "text": "..."}]: replace character spans
  tags <path>             → JSON ["TAG", ...] every {{TAG}} in order (repeats included)
  sanitize <path> <out>   → JSON {"outside": [{"where": "...", "text": "..."}]}: a copy without what a
                            client's Word carries outside its text (tracked changes accepted, comments,
                            links and content controls flattened, author/title cleared), plus the text
                            of headers, footers and notes, to check for a client's data

Edits keep each paragraph's style: the new text goes into the first run (its font/bold/size),
inserted paragraphs copy the style of the paragraph they follow. Spans and tag filling work inside
the runs, so the rest of the paragraph keeps its own bold/italics and each value its run's format.
"""
import copy
import json
import re
import sys

from docx import Document
from docx.oxml.ns import qn

TAG = re.compile(r"\{\{([^}]+)\}\}")


def paragraphs(doc):
    """Every paragraph of the body in document order (table cells where the table is).
    The fallback copy of a text box (mc:Fallback) is skipped so nothing is listed twice."""
    from docx.text.paragraph import Paragraph
    fallback = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"
    out = []
    for p in doc.element.body.iter(qn("w:p")):
        if any(a.tag == fallback for a in p.iterancestors()):
            continue
        out.append(Paragraph(p, doc._body))
    return out


def set_text(p, text):
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(text)


def text_of(p):
    return "".join(r.text for r in p.runs)


FIELD_PARTS = (qn("w:fldChar"), qn("w:instrText"))


def touches_field(p, start, end):
    pos = 0
    for r in p.runs:
        rs, re_ = pos, pos + len(r.text)
        pos = re_
        inside = rs < end and re_ > start if re_ > rs else start < rs < end
        if inside and any(c.tag in FIELD_PARTS for c in r._r):
            return True
    return False


def replace_spans(p, spans):
    """spans: [(start, end, text)] over text_of(p). Overlapping or out-of-range spans are skipped,
    and so are spans over part of a Word field (its markers would be lost)."""
    length = len(text_of(p))
    ok = []
    for start, end, text in spans:  # first come first served
        if (0 <= start < end <= length and all(end <= s or start >= e for s, e, _ in ok)
                and not touches_field(p, start, end)):
            ok.append((start, end, text))
    for start, end, text in sorted(ok, reverse=True):  # from the end, so earlier offsets stay valid
        pos, first = 0, True
        for r in p.runs:
            t = r.text
            rs, re_ = pos, pos + len(t)
            pos = re_
            if re_ <= start or rs >= end:
                continue
            a, b = max(start, rs) - rs, min(end, re_) - rs
            r.text = t[:a] + (text if first else "") + t[b:]
            first = False


def cmd_list(path):
    doc = Document(path)
    print(json.dumps([{"i": i, "text": text_of(p)} for i, p in enumerate(paragraphs(doc))], ensure_ascii=False))


def cmd_tags(path):
    doc = Document(path)
    print(json.dumps([m.group(1).strip() for p in paragraphs(doc) for m in TAG.finditer(text_of(p))], ensure_ascii=False))


def cmd_spans(path, out, items):
    doc = Document(path)
    ps = paragraphs(doc)
    by_p = {}
    for it in items:
        i = int(it.get("i", -1))
        if 0 <= i < len(ps):
            by_p.setdefault(i, []).append((int(it["start"]), int(it["end"]), str(it.get("text", ""))))
    for i, spans in by_p.items():
        replace_spans(ps[i], spans)
    doc.save(out)
    print(out)


def cmd_hastags(path):
    doc = Document(path)
    print("1" if any(TAG.search(p.text) for p in paragraphs(doc)) else "0")


def cmd_apply(path, out, ops):
    doc = Document(path)
    ps = paragraphs(doc)
    # inserts after the same paragraph keep their order: add them in reverse
    for op in reversed([o for o in ops if o.get("op") == "insert_after"]):
        i = int(op["i"])
        if 0 <= i < len(ps):
            new = copy.deepcopy(ps[i]._p)
            ps[i]._p.addnext(new)
            from docx.text.paragraph import Paragraph
            set_text(Paragraph(new, ps[i]._parent), str(op.get("text", "")))
    for op in ops:
        i = int(op.get("i", -1))
        if not 0 <= i < len(ps):
            continue
        if op.get("op") == "replace":
            set_text(ps[i], str(op.get("text", "")))
        elif op.get("op") == "delete":
            el = ps[i]._p
            el.getparent().remove(el)
    doc.save(out)
    print(out)


def cmd_fill(path, out, values):
    """A value can be a list: one per occurrence of the tag, in document order."""
    doc = Document(path)
    seen = {}
    for p in paragraphs(doc):
        spans = []
        for m in TAG.finditer(text_of(p)):
            key = m.group(1).strip()
            n = seen.get(key, 0)
            seen[key] = n + 1
            v = values.get(key)
            if isinstance(v, list):
                v = v[n] if n < len(v) else None
            if v is not None:
                spans.append((m.start(), m.end(), str(v)))
        if spans:
            replace_spans(p, spans)
    doc.save(out)
    print(out)


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
HYPERLINK_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"
OUTSIDE = (("header", "el encabezado"), ("footer", "el pie de página"), ("footnotes", "las notas al pie"), ("endnotes", "las notas finales"))


def w(tag):
    return "{%s}%s" % (W_NS, tag)


def unwrap(el):
    """Puts the element's children where it is and drops the element."""
    parent = el.getparent()
    i = parent.index(el)
    for child in list(el):
        parent.insert(i, child)
        i += 1
    parent.remove(el)


def clean_part(root):
    for tag in ("del", "moveFrom", "commentRangeStart", "commentRangeEnd", "rPrChange", "pPrChange",
                "sectPrChange", "tblPrChange", "trPrChange", "tcPrChange", "numberingChange"):
        for el in list(root.iter(w(tag))):
            el.getparent().remove(el)
    for ref in list(root.iter(w("commentReference"))):
        run = ref.getparent()
        run.getparent().remove(run)
    for tag in ("sdtPr", "sdtEndPr"):
        for el in list(root.iter(w(tag))):
            el.getparent().remove(el)
    # deepest first, so nested wrappers all open up
    for tag in ("ins", "moveTo", "hyperlink", "smartTag", "customXml", "fldSimple", "sdtContent", "sdt"):
        for el in reversed(list(root.iter(w(tag)))):
            unwrap(el)


def cmd_sanitize(path, out):
    import zipfile
    from lxml import etree
    outside = []
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            name = item.filename
            if name.startswith("word/") and name.endswith(".xml") and "/" not in name[5:]:
                root = etree.fromstring(data)
                if name.startswith("word/comments"):
                    for child in list(root):
                        root.remove(child)
                else:
                    clean_part(root)
                    for key, where in OUTSIDE:
                        if name[5:].startswith(key):
                            for para in root.iter(w("p")):
                                text = "".join(t.text or "" for t in para.iter(w("t")))
                                if text.strip():
                                    outside.append({"where": where, "text": text})
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            elif name.endswith(".rels"):
                root = etree.fromstring(data)
                for rel in list(root):
                    if rel.get("Type") == HYPERLINK_REL:
                        root.remove(rel)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            elif name in ("docProps/core.xml", "docProps/app.xml"):
                root = etree.fromstring(data)
                keep = ("created", "modified", "Application", "AppVersion", "DocSecurity", "Template")
                for child in root:
                    if etree.QName(child).localname not in keep and len(child) == 0:
                        child.text = ""
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            dst.writestr(item, data)
    print(json.dumps({"outside": outside}, ensure_ascii=False))


if __name__ == "__main__":
    cmd, args = sys.argv[1], sys.argv[2:]
    if cmd == "list":
        cmd_list(args[0])
    elif cmd == "hastags":
        cmd_hastags(args[0])
    elif cmd == "apply":
        cmd_apply(args[0], args[1], json.load(sys.stdin))
    elif cmd == "tags":
        cmd_tags(args[0])
    elif cmd == "spans":
        cmd_spans(args[0], args[1], json.load(sys.stdin))
    elif cmd == "fill":
        cmd_fill(args[0], args[1], json.load(sys.stdin))
    elif cmd == "sanitize":
        cmd_sanitize(args[0], args[1])
    else:
        sys.exit(f"unknown command {cmd}")

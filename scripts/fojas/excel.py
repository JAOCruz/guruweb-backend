#!/usr/bin/env python3
"""Excel para que Leandro revise las fojas. Uso: excel.py <fojas.json> <salida.xlsx>"""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TIPOS = {
    "carta_simple": "Carta simple (150)",
    "carta_motivada": "Carta motivada (300)",
    "legal_simple": "Legal simple (250)",
    "legal_motivada": "Legal motivada (350)",
}
HEADERS = ["Documento", "Categoría", "Tipo", "Fojas", "De dónde salen las fojas",
           "Precio anterior (digitación)", "Precio nuevo (digitación)", "Diferencia"]
WIDTHS = [50, 32, 24, 8, 44, 18, 18, 14]


def main(src, dst):
    with open(src, encoding="utf-8") as f:
        items = json.load(f)
    wb = Workbook()
    ws = wb.active
    ws.title = "Fojas"
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, it in enumerate(items, start=2):
        ws.append([
            it.get("nombre"), it.get("categoria"),
            TIPOS.get(it.get("foja_tipo"), it.get("foja_tipo")),
            it.get("fojas"), it.get("fuente"),
            it.get("precio_anterior"), it.get("precio_nuevo"),
            f'=IF(F{i}="","",G{i}-F{i})',
        ])
        for col in (6, 7, 8):
            ws.cell(row=i, column=col).number_format = "#,##0"
    for n, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(n)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{max(ws.max_row, 1)}"
    wb.save(dst)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit("Uso: excel.py <fojas.json> <salida.xlsx>")
    main(sys.argv[1], sys.argv[2])
